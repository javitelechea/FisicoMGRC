#!/usr/bin/env python3
"""
Extract all physical test data from MGRC xlsx/docx files into a structured JSON database.
Handles: YoYo test, CMJ, Flexibility, Strength (4MR), Biometrics, Attendance, Injuries.
"""

import json
import os
import re
import sys
from datetime import datetime

import openpyxl
from docx import Document

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data_raw", "2026")
OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "..", "public", "data", "players.json")

CATEGORY_MAP = {
    "TEST 5 A": {"division": "5ta", "sub": "A", "label": "5ta A"},
    "TEST 5 B": {"division": "5ta", "sub": "B", "label": "5ta B"},
    "TEST 5 C": {"division": "5ta", "sub": "C", "label": "5ta C"},
    "TEST 5 D": {"division": "5ta", "sub": "D", "label": "5ta D"},
    "Test  6  A": {"division": "6ta", "sub": "A", "label": "6ta A"},
    "Test  6  B": {"division": "6ta", "sub": "B", "label": "6ta B"},
    "Test  6  C": {"division": "6ta", "sub": "C", "label": "6ta C"},
    "Test  6  D": {"division": "6ta", "sub": "D", "label": "6ta D"},
    "Test 7 A": {"division": "7ma", "sub": "A", "label": "7ma A"},
    "Test 7 B": {"division": "7ma", "sub": "B", "label": "7ma B"},
    "Test 7 C": {"division": "7ma", "sub": "C", "label": "7ma C"},
    "Test 7 D": {"division": "7ma", "sub": "D", "label": "7ma D"},
    "yoyo Marzo": {"division": "Sub14", "sub": "Marzo", "label": "Sub14 (Marzo)"},
}


def normalize_name(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r"\s+", " ", name)
    parts = name.split()
    if len(parts) >= 2:
        if parts[0].istitle() and parts[-1].istitle():
            return name
        return " ".join(p.capitalize() for p in parts)
    return name.capitalize()


def detect_category_from_filename(filename):
    for key, val in CATEGORY_MAP.items():
        if key.lower().replace(" ", "") in filename.lower().replace(" ", ""):
            return val
    return None


def safe_float(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", ".")
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def safe_str(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def extract_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if s:
        return s
    return None


def extract_test_sheet(ws, category_info):
    """Extract player data from a Test sheet."""
    players = []

    yoyo_date_col = None
    for row in ws.iter_rows(min_row=3, max_row=5, min_col=6, max_col=20, values_only=False):
        for cell in row:
            if cell.value and "Fecha" in str(cell.value):
                yoyo_date_col = cell.column
                break
        if yoyo_date_col:
            break

    date_cell_val = None
    if yoyo_date_col:
        for r in range(3, 6):
            v = ws.cell(row=r, column=yoyo_date_col + 1).value
            if v:
                date_cell_val = extract_date(v)
                break

    name_col = 3  # Column C
    start_row = 6

    for row_idx in range(start_row, min(ws.max_row + 1, 60)):
        name = ws.cell(row=row_idx, column=name_col).value
        if not name or not str(name).strip():
            continue

        name = normalize_name(name)
        age = safe_float(ws.cell(row=row_idx, column=4).value)
        position = safe_str(ws.cell(row=row_idx, column=5).value)

        yoyo_level = safe_float(ws.cell(row=row_idx, column=6).value)
        yoyo_meters = safe_float(ws.cell(row=row_idx, column=7).value)

        if yoyo_level is not None and (yoyo_level < 5 or yoyo_level > 25):
            yoyo_level = None
        if yoyo_meters is not None and (yoyo_meters < 40 or yoyo_meters > 3000):
            yoyo_meters = None

        player = {
            "name": name,
            "category": category_info["label"],
            "division": category_info["division"],
            "sub": category_info["sub"],
            "age": age,
            "position": position,
            "tests": [],
        }

        test_entry = {
            "date": date_cell_val or "2026-03",
            "source": "planilla",
            "yoyo": {},
            "cmj": {},
            "flexibility": {},
            "strength": {},
        }

        if yoyo_level is not None:
            test_entry["yoyo"]["level"] = yoyo_level
        if yoyo_meters is not None:
            test_entry["yoyo"]["meters"] = yoyo_meters

        # Find CMJ section - look for this player's name in the CMJ columns
        cmj_start_col = None
        for col in range(30, 90):
            header = ws.cell(row=2, column=col).value
            if header and "CMJ" in str(header).upper():
                cmj_start_col = col
                break

        if cmj_start_col:
            for search_row in range(start_row - 1, min(ws.max_row + 1, 60)):
                cmj_name = ws.cell(row=search_row, column=cmj_start_col - 2).value
                if cmj_name and normalize_name(cmj_name) == name:
                    height_col = cmj_start_col - 2 + 4
                    height = safe_float(ws.cell(row=search_row, column=height_col).value)
                    if height is not None:
                        test_entry["cmj"]["height"] = height
                    break

        # Find Flexibility section
        flex_start_col = None
        for col in range(40, 90):
            header = ws.cell(row=2, column=col).value
            if header and "Flexibilidad" in str(header):
                flex_start_col = col
                break

        if flex_start_col:
            flex_labels = ["isquiotibial", "psoas", "aductor", "cuadriceps", "gluteos", "gemelos"]
            for search_row in range(start_row - 1, min(ws.max_row + 1, 60)):
                flex_name = ws.cell(row=search_row, column=flex_start_col).value
                if flex_name and normalize_name(flex_name) == name:
                    for li, label in enumerate(flex_labels):
                        flex_header_col = flex_start_col + 4 + li
                        val_str = safe_str(ws.cell(row=search_row, column=flex_header_col).value)
                        if val_str:
                            test_entry["flexibility"][label] = val_str
                    break

        # Find Strength section
        str_start_col = None
        for col in range(50, 100):
            header = ws.cell(row=2, column=col).value
            if header and "Fuerza" in str(header):
                str_start_col = col
                break

        if str_start_col:
            strength_tests = {
                "press_pecho": 1,
                "dorsal_remo": 4,
                "sentadilla": 7,
                "hamstring": 10,
                "pivot_press": 13,
            }
            for search_row in range(start_row - 1, min(ws.max_row + 1, 60)):
                str_name = ws.cell(row=search_row, column=str_start_col).value
                if str_name and normalize_name(str_name) == name:
                    peso = safe_float(ws.cell(row=search_row, column=str_start_col + 3).value)
                    if peso:
                        test_entry["strength"]["peso"] = peso
                    for test_name, offset in strength_tests.items():
                        val = safe_float(
                            ws.cell(row=search_row, column=str_start_col + 4 + offset).value
                        )
                        if val and val != 0:
                            test_entry["strength"][test_name] = val
                    break

        has_data = (
            test_entry["yoyo"]
            or test_entry["cmj"]
            or test_entry["flexibility"]
            or test_entry["strength"]
        )
        if has_data:
            player["tests"].append(test_entry)

        players.append(player)

    return players


def extract_docx_history(filepath):
    """Extract player history from a DOCX file."""
    doc = Document(filepath)

    player_name = None
    for para in doc.paragraphs:
        text = para.text.strip()
        if "Jugadora:" in text or "jugadora:" in text.lower():
            player_name = text.split(":", 1)[1].strip()
            break

    if not player_name:
        basename = os.path.splitext(os.path.basename(filepath))[0]
        basename = re.sub(r"^Historia\s+[Ff]isica\s+", "", basename)
        player_name = basename

    player_name = normalize_name(player_name)

    history = {
        "name": player_name,
        "biometrics": [],
        "yoyo": [],
        "cmj": [],
        "strength": [],
        "attendance": [],
        "injuries": [],
    }

    months = ["Feb.", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Sept", "Oct.", "Nov."]

    for table in doc.tables:
        if len(table.rows) < 2:
            continue

        header = [cell.text.strip() for cell in table.rows[0].cells]

        for row_idx in range(1, len(table.rows)):
            row_cells = [cell.text.strip() for cell in table.rows[row_idx].cells]

            year = row_cells[0] if row_cells[0] else None
            division = row_cells[1].replace("\n", " ") if len(row_cells) > 1 else None
            pf = row_cells[2] if len(row_cells) > 2 else None
            metric = row_cells[3] if len(row_cells) > 3 else None

            if not metric:
                continue

            if "Biometricas" in header or "Biometricas" in str(header):
                for mi, month in enumerate(months):
                    col_idx = 4 + mi
                    if col_idx < len(row_cells) and row_cells[col_idx]:
                        val = safe_float(row_cells[col_idx])
                        if val is not None:
                            history["biometrics"].append(
                                {
                                    "year": year,
                                    "division": division,
                                    "metric": metric,
                                    "month": month,
                                    "value": val,
                                }
                            )

            elif "YY" in str(header) or "Nivel" in metric or "Metros" in metric:
                for mi, month in enumerate(months):
                    col_idx = 4 + mi
                    if col_idx < len(row_cells) and row_cells[col_idx]:
                        val = safe_float(row_cells[col_idx])
                        if val is not None:
                            history["yoyo"].append(
                                {
                                    "year": year,
                                    "division": division,
                                    "metric": metric,
                                    "month": month,
                                    "value": val,
                                }
                            )

            elif "CMJ" in str(header) or "Altura" in metric:
                for mi, month in enumerate(months):
                    col_idx = 4 + mi
                    if col_idx < len(row_cells) and row_cells[col_idx]:
                        val = safe_float(row_cells[col_idx])
                        if val is not None:
                            history["cmj"].append(
                                {
                                    "year": year,
                                    "division": division,
                                    "metric": metric,
                                    "month": month,
                                    "value": val,
                                }
                            )

            elif "Fuerza" in str(header) or metric in [
                "Sent.",
                "H. T",
                "Dorsal",
                "Pectoral",
            ]:
                for mi, month in enumerate(months):
                    col_offset = 6 if len(header) > 14 else 4
                    col_idx = col_offset + mi
                    if col_idx < len(row_cells) and row_cells[col_idx]:
                        val = safe_float(row_cells[col_idx])
                        if val is not None:
                            history["strength"].append(
                                {
                                    "year": year,
                                    "division": division,
                                    "exercise": metric,
                                    "month": month,
                                    "value": val,
                                }
                            )

            elif "Presentismo" in str(header) or "%" in metric:
                for mi, month in enumerate(months):
                    col_idx = 4 + mi
                    if col_idx < len(row_cells) and row_cells[col_idx]:
                        val = safe_float(row_cells[col_idx])
                        if val is not None:
                            history["attendance"].append(
                                {
                                    "year": year,
                                    "division": division,
                                    "month": month,
                                    "percentage": val,
                                }
                            )

            elif "Lesion" in str(header) or "Caracteristica" in metric:
                for mi, month in enumerate(months):
                    col_idx = 4 + mi
                    if col_idx < len(row_cells) and row_cells[col_idx]:
                        history["injuries"].append(
                            {
                                "year": year,
                                "division": division,
                                "month": month,
                                "description": row_cells[col_idx],
                            }
                        )

    return history


def merge_player_data(xlsx_players, docx_histories):
    """Merge xlsx test data with docx history, using fuzzy name matching."""
    merged = {}

    for p in xlsx_players:
        key = p["name"].lower().strip()
        if key not in merged:
            merged[key] = {
                "id": len(merged) + 1,
                "name": p["name"],
                "category": p["category"],
                "division": p["division"],
                "sub": p.get("sub"),
                "age": p.get("age"),
                "position": p.get("position"),
                "tests": p.get("tests", []),
                "history": None,
                "name_variants": [p["name"]],
            }
        else:
            if p.get("tests"):
                merged[key]["tests"].extend(p["tests"])
            if p.get("age") and not merged[key]["age"]:
                merged[key]["age"] = p["age"]
            if p.get("position") and not merged[key]["position"]:
                merged[key]["position"] = p["position"]
            if p["name"] not in merged[key]["name_variants"]:
                merged[key]["name_variants"].append(p["name"])

    def fuzzy_match(name, candidates, threshold=0.7):
        name_lower = name.lower().strip()
        name_parts = set(name_lower.split())

        best_match = None
        best_score = 0

        for cand in candidates:
            cand_lower = cand.lower().strip()
            if name_lower == cand_lower:
                return cand

            cand_parts = set(cand_lower.split())
            if name_parts & cand_parts:
                intersection = len(name_parts & cand_parts)
                union = len(name_parts | cand_parts)
                score = intersection / union
                if score > best_score and score >= threshold:
                    best_score = score
                    best_match = cand

            if not best_match:
                from difflib import SequenceMatcher

                ratio = SequenceMatcher(None, name_lower, cand_lower).ratio()
                if ratio > best_score and ratio >= threshold:
                    best_score = ratio
                    best_match = cand

        return best_match

    for hist in docx_histories:
        hist_name_lower = hist["name"].lower().strip()

        if hist_name_lower in merged:
            merged[hist_name_lower]["history"] = hist
            continue

        match = fuzzy_match(hist["name"], list(merged.keys()))
        if match:
            merged[match]["history"] = hist
            if hist["name"] not in merged[match]["name_variants"]:
                merged[match]["name_variants"].append(hist["name"])
        else:
            key = hist_name_lower
            merged[key] = {
                "id": len(merged) + 1,
                "name": hist["name"],
                "category": None,
                "division": None,
                "sub": None,
                "age": None,
                "position": None,
                "tests": [],
                "history": hist,
                "name_variants": [hist["name"]],
            }

    result = sorted(merged.values(), key=lambda x: x["name"])
    for i, p in enumerate(result):
        p["id"] = i + 1
    return result


def main():
    xlsx_players = []

    for filename in sorted(os.listdir(DATA_DIR)):
        if not filename.endswith(".xlsx"):
            continue

        filepath = os.path.join(DATA_DIR, filename)
        cat = detect_category_from_filename(filename)
        if not cat:
            print(f"  [SKIP] No category match for: {filename}")
            continue

        print(f"  Processing: {filename} -> {cat['label']}")

        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            test_sheet = None
            for sn in wb.sheetnames:
                if sn.lower() in ("test", "tests"):
                    test_sheet = wb[sn]
                    break

            if test_sheet:
                players = extract_test_sheet(test_sheet, cat)
                xlsx_players.extend(players)
                print(f"    -> {len(players)} players extracted")
            else:
                print(f"    -> No 'Test' sheet found (sheets: {wb.sheetnames})")
        except Exception as e:
            print(f"    ERROR: {e}")

    print(f"\nTotal players from xlsx: {len(xlsx_players)}")

    docx_dir = os.path.join(DATA_DIR, "Historias Fisicas")
    docx_histories = []

    if os.path.isdir(docx_dir):
        for filename in sorted(os.listdir(docx_dir)):
            if not filename.endswith(".docx"):
                continue
            filepath = os.path.join(docx_dir, filename)
            print(f"  Processing DOCX: {filename}")
            try:
                hist = extract_docx_history(filepath)
                docx_histories.append(hist)
            except Exception as e:
                print(f"    ERROR: {e}")

    print(f"\nTotal histories from docx: {len(docx_histories)}")

    all_players = merge_player_data(xlsx_players, docx_histories)
    print(f"\nMerged unique players: {len(all_players)}")

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)

    categories = {}
    for p in all_players:
        cat = p.get("category") or "Sin categoría"
        if cat not in categories:
            categories[cat] = []
        categories[cat].append(p["name"])

    output = {
        "generated_at": datetime.now().isoformat(),
        "total_players": len(all_players),
        "categories": categories,
        "players": all_players,
    }

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2, default=str)

    print(f"\nData saved to: {OUTPUT_FILE}")
    print(f"Categories found: {list(categories.keys())}")
    for cat, names in categories.items():
        print(f"  {cat}: {len(names)} players")


if __name__ == "__main__":
    main()
